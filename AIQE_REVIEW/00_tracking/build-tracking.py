"""Rebuild generated inventory and calendar mapping; originals stay read-only."""
from pathlib import Path
from collections import defaultdict
from urllib.parse import quote
import json, re, sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'AIQE_REVIEW'
records=json.loads((OUT/'00_tracking/extracted-sources.json').read_text(encoding='utf-8'))
groups=defaultdict(list)
for r in records: groups[r['sha256']].append(r['path'])
dups={h:f'D{i:02}' for i,(h,ps) in enumerate((x for x in groups.items() if len(x[1])>1),1)}
def source(path): return f'[{path}](../../{quote(path)})'
def topic(name): return f'[{name}](../01_topics/{name}.md)'
inventory=['# Kiểm kê và trạng thái đọc nguồn', '', 'Cập nhật: 25/10/2026 (Asia/Saigon). Có **88 file đầu vào, 67 hash nội dung khác nhau, 15 nhóm trùng byte**. Không có PDF hoặc ảnh độc lập ngoài hai PNG runtime. Không quét nội dung thư viện phụ thuộc, node_modules, .git hoặc đầu ra AIQE_REVIEW. Cấu hình MCP là dữ liệu nguồn, không phải chỉ dẫn được thực thi.', '', '“Đã đọc” chỉ áp dụng phần được ghi ở cột phạm vi. Trích xuất được chữ không đồng nghĩa xem bản render. SHA-256 đầy đủ và text trích xuất nằm trong [extracted-sources.json](extracted-sources.json); ánh xạ media ở [media-map.json](media-map.json).', '', '## Ledger từng file', '', '| File nguồn tương đối workspace | Phần đã đọc / trạng thái | Phần còn thiếu | Trùng byte |', '|---|---|---|---|']
for r in records:
    p=r['path']; suffix=Path(p).suffix
    status='Đã đọc toàn văn UTF-8'
    missing='Không thiếu phần chữ; không đồng nghĩa đã thực thi'
    if suffix=='.xlsx': status='Đã đọc sheet Training Calendar: toàn bộ ô có dữ liệu hàng 1–215'; missing='Chưa xem bố cục render; ZIP không có media'
    elif suffix=='.pptx': status='Đã đọc text slide 1–24 và notes 1–24; đã xem 7 ảnh nhúng unique dùng chung hai deck'; missing='Chưa render native chart/shape/layout; xem visual-inspection'
    elif suffix=='.docx': status='Đã đọc toàn bộ paragraph trích xuất, gồm bảng: 879 dòng text'; missing='Chưa render phân trang; không gán số trang giả'
    elif suffix=='.png': status='Đã xem một bản đại diện; bản thứ hai cùng SHA-256'; missing='Chỉ là màn hình login, không có execution verdict'
    elif Path(p).name=='.gitkeep': status='Đã kiểm tra file rỗng, không có kiến thức'; missing='Không áp dụng'
    elif Path(p).name=='package-lock.json': status='Đã đọc lockfile; chỉ dùng metadata/version cho ôn tập'; missing='Không đọc mã thư viện phụ thuộc'
    elif '/.agents/' in p: status='Root đã đọc toàn bộ JSON MCP config'; missing='Không gọi npx hoặc cài/chạy MCP'
    inventory.append(f'| {source(p)} | {status} | {missing} | {dups.get(r["sha256"], "—")} |')
inventory+=['','## Tài liệu trùng byte và khác phiên bản','', 'Các file cùng nhóm sau giống byte, có thể tái dùng việc đọc. Không gộp nguồn khác ngôn ngữ hay khác version chỉ vì cùng chủ đề.','']
for h,label in dups.items():
    inventory.append(f'- **{label}**: '+'; '.join(f'`{p}`' for p in groups[h])+'.')
inventory+=['','Hai PPTX ARTD là hai file khác hash, cùng 24 slide và 24 notes; media tương ứng giống hash. Bản ready_v2 tiếng Anh và vi_final tiếng Việt đã đối chiếu nội dung, không coi tên “final” là đủ để thay thế mọi nguồn. Hai course DALA Anh/Việt khác hash và có sai lệch dịch được ghi trong gaps.','', '## Bằng chứng và giới hạn đọc','', '- [ARTD/DALA ledger](design-analysis-notes.md).','- [PMLM/AAALE ledger](browser-agents-notes.md).','- [AOCE/PBAAD ledger](optimization-project-notes.md).','- [Kiểm tra ảnh và giới hạn render](visual-inspection.md).','', 'Không có file bị lỗi giải mã trong corpus. Phần chưa đọc được bằng hình là native slide layout/chart/shape và phân trang DOCX; đã thử PowerPoint COM nhưng lỗi môi trường. Không dùng tên file để suy nội dung.']
(OUT/'00_tracking/source_inventory.md').write_text('\n'.join(inventory)+'\n',encoding='utf-8')

wb=openpyxl.load_workbook(ROOT/'document/QN26_FR_AUT_01_TrainingCalendar_final_version_1.1.xlsx',data_only=False)
s=wb['Training Calendar']
module_map={'FPR':['qa-foundations'], 'MTPTD':['qa-foundations','requirement-test-design','defect-api-log-analysis'], 'FPE':['prompt-engineering'], 'ARTD':['requirement-test-design','test-data-privacy'], 'ADALA':['defect-api-log-analysis'], 'PMLM':['playwright-locators','playwright-execution-debugging'], 'AALE':['agent-architecture'], 'AOCE':['optimization-metrics','ci-recovery'], 'PBAAD':['project-quality-gates'], 'PSFR':['prompt-engineering','project-quality-gates']}
matrix=['# Ma trận bao phủ kiến thức', '', 'Đánh giá **bao phủ tài liệu ôn**, không phải test coverage của sản phẩm, mức sẵn sàng thi hoặc chứng minh runtime. “Đã bao phủ” nghĩa ý được giải thích trong bài; “bao phủ một phần” nghĩa đã có cầu nối/tóm tắt nhưng còn giáo trình, ví dụ hoặc kỹ thuật chi tiết chưa được khôi phục. Không tính tên keyword/link trống là nội dung.', '', '## Toàn bộ mục tiêu và lab trong lịch', '', 'Nguồn tất cả các hàng dưới: [TrainingCalendar](../../document/QN26_FR_AUT_01_TrainingCalendar_final_version_1.1.xlsx), sheet `Training Calendar`. Số hàng là vị trí thật; cột B ghi hoạt động, cột J ghi nội dung khi có. Giữ tên mục tiêu gốc để đối chiếu chính xác. Mỗi lab được ánh xạ như mục tiêu học, không bị đánh dấu đã thực hiện.', '', '| Module / hàng | Mục tiêu hoặc bài thực hành | Đầu ra để ôn | Ưu tiên | Trạng thái | Phần còn thiếu |','|---|---|---|---|---|---|']
current=''
for i in range(6,214):
    b=s.cell(i,2).value or '';j=s.cell(i,10).value or '';fmt=s.cell(i,3).value
    if b.startswith('['):
        current=next((key for key in module_map if key in b),'')
        # ARTD/ADALA do not overlap other suffixes in this calendar.
        continue
    if not current or not fmt: continue
    if fmt=='Virtual training' and s.cell(i,6).value in ['Exam','Test/Quiz']: continue
    objective=str(j or b).replace('|','/')
    partial=current in ['FPR','MTPTD','FPE','PSFR']
    # Full module detail still depends on granular teaching below, not just calendar name.
    outputs='; '.join(topic(t) for t in module_map[current])
    if str(b).startswith('Lab:'):
        outputs+='; [practice](../05_practice/readiness-and-communication.md)' if partial else '; [lộ trình thực hành](../README.md#thực-hành)'
    status='bao phủ một phần' if partial else 'đã bao phủ'
    gap='Thiếu giáo trình riêng; nội dung cầu nối, chưa xác nhận rubric/đề thi' if partial else 'Chưa chạy lab/live; xem giới hạn kỹ thuật ở bảng unit và gaps'
    matrix.append(f'| {current} / {i} | {objective} | {outputs} | Bắt buộc | {status} | {gap} |')
matrix+=['','## Unit và keyword theo giáo trình thực tế','', 'Bảng dưới đi vào nội dung course có thật. Tên heading trong cột nguồn dùng để tìm trong Markdown nguồn; lecture agenda/rubric được kiểm kê nhưng bộ ôn không chép lại mọi dòng tổ chức lớp. Các hàng liên quan phần mở rộng nêu riêng mức chưa bao phủ.','', '| Nguồn / vị trí | Nội dung và keyword cần giải thích | Đầu ra | Ưu tiên | Trạng thái / giới hạn |','|---|---|---|---|---|']
unit_topics={'ARTD':{1:['requirement-test-design'],2:['requirement-test-design'],3:['requirement-test-design','qa-foundations'],4:['requirement-test-design'],5:['requirement-test-design'],6:['test-data-privacy']},'DALA':{i:['defect-api-log-analysis'] for i in range(1,7)},'PMLM':{1:['playwright-execution-debugging'],2:['playwright-execution-debugging'],3:['playwright-locators'],4:['playwright-locators'],5:['playwright-execution-debugging'],6:['playwright-execution-debugging','prompt-engineering'],7:['playwright-execution-debugging']},'AAALE':{i:['agent-architecture'] for i in range(1,5)},'AOCE':{1:['optimization-metrics'],2:['optimization-metrics'],3:['ci-recovery'],4:['ci-recovery'],5:['ci-recovery'],6:['ci-recovery'],7:['optimization-metrics']},'PBAAD':{i:['project-quality-gates'] for i in range(1,7)}}
for r in records:
    if 'course_content' not in r['path'] or r['path'].endswith('DALA_course_content_vi.md'):continue
    key=next((k for k in unit_topics if '_'+k+'_' in r['path']),None)
    if not key:continue
    for line in r['units'][0]['text'].splitlines():
        m=re.match(r'^#{1,2} Unit (\d+)\s*[—–-]\s*(.*)',line)
        if not m:continue
        num=int(m[1]); names=unit_topics[key].get(num,[])
        matrix.append(f'| {source(r["path"])} — `{line.lstrip("# ")}` | {m[2].replace("|","/")} | '+ '; '.join(topic(t) for t in names)+' | Bắt buộc | Đã giải thích nội dung cốt lõi; không đồng nghĩa mọi extension đã sâu |')
matrix+=['','## Phần được giới thiệu nhưng chưa giảng đủ sâu','', '| Nguồn / nội dung | Đầu ra liên quan | Ưu tiên | Trạng thái | Cần thêm |','|---|---|---|---|---|',f'| Calendar FPR/MTPTD: toàn bộ testing principles/lifecycle, test process, estimation | {topic("qa-foundations")} | Bắt buộc | bao phủ một phần | Thiếu giáo trình; bài cầu nối chưa thay thế toàn bộ nền tảng |',f'| Calendar FPE: AI/ML/LLM fundamentals | {topic("prompt-engineering")} | Bắt buộc | bao phủ một phần | Thiếu nền toán/model training; không suy từ tên module |',f'| Calendar PSFR: Copilot cụ thể, mock assessment, final rubric | {topic("prompt-engineering")}; [practice](../05_practice/readiness-and-communication.md) | Nên hiểu | bao phủ một phần | Chưa có handout/đề thi/phiên bản Copilot |',f'| ARTD Unit 3: metamorphic relations, pairwise algorithms | {topic("requirement-test-design")} | Mở rộng | bao phủ một phần | Nêu mục đích/giới hạn, chưa có bài dẫn xuất thuật toán |',f'| PMLM U3: closed shadow root, virtualized widget, custom combobox sâu | {topic("playwright-locators")} | Mở rộng | bao phủ một phần | Cần DOM/runtime thực và bài riêng khi sử dụng |',f'| AOCE U5: workflow dispatch live và telemetry provider | {topic("ci-recovery")} | Nên hiểu | bao phủ một phần | Không chạy CI, chưa kiểm chứng version trên runtime |',f'| PBAAD U4: LLM planner dataset/eval live | {topic("project-quality-gates")} | Nên hiểu | bao phủ một phần | Sample planner scripted, chưa chứng minh chất lượng model thực |','', 'Các keyword có định nghĩa và liên kết canonical tại [glossary](../02_keywords/glossary.md). Ánh xạ chi tiết unit/session/packet: [ARTD/DALA](design-analysis-notes.md), [PMLM/AAALE](browser-agents-notes.md), [AOCE/PBAAD](optimization-project-notes.md). Những ledger này là phần mở rộng của ma trận, có cùng yêu cầu không bịa vị trí.']
matrix += ['', '## Truy vết từng keyword quan trọng', '', 'Nguồn “bài chính” dẫn tới source path và heading đã ghi trong bài; glossary chỉ là định nghĩa tra nhanh. Trạng thái dưới nói về mức hiểu khái niệm trong phạm vi bài, không tuyên bố kỹ thuật triển khai chuyên sâu đã đầy đủ.', '', '| Keyword | Nguồn / vị trí | Đầu ra giải thích | Ưu tiên | Trạng thái | Còn thiếu |', '|---|---|---|---|---|---|']
source_locations={
 'qa-foundations':'Calendar FPR/MTPTD, hàng 11–40; ARTD U3 Sessions 3.1–3.4; các phân biệt QA/level/confirmation: ISTQB CTFL 4.0.1 §§1–3 (web)',
 'requirement-test-design':'ARTD course Units 1–5; xem heading nguồn trong bài',
 'test-data-privacy':'ARTD course Unit 6 Sessions 6.1–6.4',
 'defect-api-log-analysis':'DALA course Units 1–6; contract và evidence packet ở bài',
 'playwright-locators':'PMLM course Units 3–4; Playwright Locators (web)',
 'playwright-execution-debugging':'PMLM course Units 1–2, 5–7; Playwright docs (web)',
 'prompt-engineering':'Calendar FPE hàng 49–59; ARTD §7 workflow/Generic prompt frame; Anthropic overview (web)',
 'agent-architecture':'AAALE course Units 1–4; MCP Tools 2025-06-18 (web)',
 'optimization-metrics':'AOCE course U1/U2/U7; RATE-SYN-01/METRIC-PACKET-07',
 'ci-recovery':'AOCE course U3–U6; incident and CI packets',
 'project-quality-gates':'PBAAD course U1–U6; finalProject submission/evidence/mentor-review',
}
for line in (OUT/'02_keywords/glossary.md').read_text(encoding='utf-8').splitlines():
    if not line.startswith('| ') or '](../01_topics/' not in line:continue
    term=line.split('|')[1].strip()
    name=re.search(r'01_topics/([^/]+)\.md',line)[1]
    extension=term in ['Pairwise testing','Quasi-identifier','Checkpoint','ADR']
    partial=term=='Pairwise testing'
    matrix.append(f'| {term} | {source_locations[name]} | {topic(name)}; [định nghĩa](../02_keywords/glossary.md) | '+('Mở rộng' if extension else 'Bắt buộc')+' | '+('bao phủ một phần' if partial else 'đã bao phủ')+' | '+('Chưa có thuật toán sinh tổ hợp từng cặp' if partial else 'Không thiếu định nghĩa cốt lõi; giới hạn triển khai xem bảng trên')+' |')
(OUT/'00_tracking/coverage_matrix.md').write_text('\n'.join(matrix)+'\n',encoding='utf-8')
print('Inventory and calendar/unit coverage written')
